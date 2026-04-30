from __future__ import annotations

import csv
import json
import queue
import random
import threading
import time
import tkinter as tk
import urllib.request
from io import BytesIO
from pathlib import Path
from tkinter import filedialog, messagebox, ttk

try:
    from openpyxl import load_workbook
except ImportError:
    load_workbook = None

try:
    from PIL import Image, ImageDraw, ImageOps, ImageTk
except ImportError:
    Image = ImageDraw = ImageOps = ImageTk = None

from .audio_player import AudioPlayer, AudioPlayerError
from .downloader import AudioDownloadError, AudioStream, BiliAudioDownloader, BiliSearchResult
from .library import FavoriteTrack, MusicLibrary, Track
from .paths import SEARCH_CACHE_FILE


SEARCH_LIMIT = 12
PROGRESS_SCALE_MAX = 1000.0
WINDOW_BG = "#111318"
SURFACE_BG = "#181b22"
ELEVATED_BG = "#202530"
CARD_BG = "#242a37"
ACCENT = "#ff5c7c"
ACCENT_HOVER = "#ff7892"
TEXT_PRIMARY = "#f7f8fb"
TEXT_SECONDARY = "#b7bfd0"
TEXT_MUTED = "#7f889b"
BORDER = "#2b3242"
TREE_BG = "#191e27"
TREE_ALT = "#1f2530"
TREE_SELECTED = "#3a2431"
IMPORT_DELAY_RANGE = (2.0, 5.0)
IMPORT_RETRY_DELAYS = (5.0, 15.0, 30.0)


class BiliMusicPlayerApp(tk.Tk):
    def __init__(self) -> None:
        super().__init__()
        self.title("Bili Music Player")
        self.geometry("1240x780")
        self.minsize(1020, 680)
        self.configure(bg=WINDOW_BG)

        self.library = MusicLibrary()
        self.downloader = BiliAudioDownloader()
        self.player = AudioPlayer()
        self.events: queue.Queue[tuple[str, object]] = queue.Queue()

        self.search_results: list[BiliSearchResult] = []
        self.is_search_running = False
        self.is_media_task_running = False
        self.is_dragging_progress = False
        self.current_duration: float | None = None
                self.playback_started_at = 0.0
        self.favorite_play_mode = "sequential"
        self.current_favorite_url: str | None = None
        self.favorite_mode_var = tk.StringVar(value="播放模式：顺序")

        self.query_var = tk.StringVar()
        self.cache_after_stream_var = tk.BooleanVar(value=False)
        self.status_var = tk.StringVar(value="准备就绪，可以搜歌了。")
        self.pause_text = tk.StringVar(value="暂停")
        self.now_playing_var = tk.StringVar(value="未播放")
        self.time_var = tk.StringVar(value="00:00 / --:--")
        self.progress_var = tk.DoubleVar(value=0.0)
        self.search_summary_var = tk.StringVar(value="搜你想听的歌、歌手、MV 或 BV 号。")
        self.library_summary_var = tk.StringVar(value="本地还没有缓存音频。")
        self.favorite_summary_var = tk.StringVar(value="把喜欢的歌收进列表，下次一点就播。")
        self.favorite_mode_summary_var = tk.StringVar(value="当前模式：顺序播放")
        self.selection_title_var = tk.StringVar(value="等待选择内容")
        self.selection_meta_var = tk.StringVar(value="双击搜索结果即可开始播放，也可以保存到本地。")
        self.selection_extra_var = tk.StringVar(value="支持在线播放和缓存到本地。")
        self.play_mode_var = tk.StringVar(value="空闲中")
        self.stats_var = tk.StringVar(value="0 首缓存 · 0 条搜索结果")
        self.transport_hint_var = tk.StringVar(value="在线播放优先，失败时自动转临时缓存。")
                self._thumbnail_request_token = 0
        self._artwork_placeholder = self._create_placeholder_artwork()

        self.search_cache = self._load_search_cache()

        self._configure_style()
        self._build_ui()
        self._load_library()
        self._load_favorites()
        self._refresh_dashboard()
        self.after(150, self._process_events)
        self.after(500, self._update_progress_ui)
        self.protocol("WM_DELETE_WINDOW", self._on_close)

    def _configure_style(self) -> None:
        style = ttk.Style(self)
        try:
            if "clam" in style.theme_names():
                style.theme_use("clam")
        except tk.TclError:
            pass

        default_font = ("Microsoft YaHei UI", 10)
        title_font = ("Microsoft YaHei UI", 12, "bold")

        style.configure("App.TFrame", background=WINDOW_BG)
        style.configure("Surface.TFrame", background=SURFACE_BG)
        style.configure("Card.TFrame", background=CARD_BG)
        style.configure("TLabel", background=WINDOW_BG, foreground=TEXT_PRIMARY, font=default_font)
        style.configure("Muted.TLabel", background=WINDOW_BG, foreground=TEXT_MUTED, font=default_font)
        style.configure("Surface.TLabel", background=SURFACE_BG, foreground=TEXT_PRIMARY, font=default_font)
        style.configure("Card.TLabel", background=CARD_BG, foreground=TEXT_PRIMARY, font=default_font)
        style.configure("HeroTitle.TLabel", background=SURFACE_BG, foreground=TEXT_PRIMARY, font=("Microsoft YaHei UI", 22, "bold"))
        style.configure("HeroSub.TLabel", background=SURFACE_BG, foreground=TEXT_SECONDARY, font=("Microsoft YaHei UI", 10))
        style.configure("SectionTitle.TLabel", background=WINDOW_BG, foreground=TEXT_PRIMARY, font=title_font)
        style.configure("CardTitle.TLabel", background=CARD_BG, foreground=TEXT_PRIMARY, font=title_font)
        style.configure("NowPlaying.TLabel", background=CARD_BG, foreground=TEXT_PRIMARY, font=("Microsoft YaHei UI", 16, "bold"))
        style.configure("Meta.TLabel", background=CARD_BG, foreground=TEXT_SECONDARY, font=("Microsoft YaHei UI", 10))
        style.configure("Status.TLabel", background=WINDOW_BG, foreground=TEXT_SECONDARY, padding=(10, 8), font=("Microsoft YaHei UI", 9))
        style.configure("Summary.TLabel", background=WINDOW_BG, foreground=TEXT_SECONDARY, font=("Microsoft YaHei UI", 9))
        style.configure("Panel.TLabelframe", background=WINDOW_BG, borderwidth=0, relief="flat")
        style.configure("Panel.TLabelframe.Label", background=WINDOW_BG, foreground=TEXT_PRIMARY, font=title_font)

        style.configure(
            "Search.TEntry",
            fieldbackground=ELEVATED_BG,
            foreground=TEXT_PRIMARY,
            bordercolor=BORDER,
            lightcolor=BORDER,
            darkcolor=BORDER,
            insertcolor=TEXT_PRIMARY,
            padding=(12, 10),
        )
        style.map("Search.TEntry", bordercolor=[("focus", ACCENT)], lightcolor=[("focus", ACCENT)], darkcolor=[("focus", ACCENT)])

        style.configure(
            "Primary.TButton",
            background=ACCENT,
            foreground="white",
            borderwidth=0,
            focusthickness=0,
            padding=(16, 10),
            font=("Microsoft YaHei UI", 10, "bold"),
        )
        style.map(
            "Primary.TButton",
            background=[("active", ACCENT_HOVER), ("disabled", "#8b4251")],
            foreground=[("disabled", "#ead5db")],
        )

        style.configure(
            "Secondary.TButton",
            background=ELEVATED_BG,
            foreground=TEXT_PRIMARY,
            borderwidth=0,
            focusthickness=0,
            padding=(14, 10),
        )
        style.map(
            "Secondary.TButton",
            background=[("active", CARD_BG), ("disabled", "#20232b")],
            foreground=[("disabled", TEXT_MUTED)],
        )

        style.configure(
            "Horizontal.Player.TScale",
            background=CARD_BG,
            troughcolor=ELEVATED_BG,
            sliderthickness=16,
        )

        style.configure(
            "Accent.TCheckbutton",
            background=SURFACE_BG,
            foreground=TEXT_SECONDARY,
            indicatorcolor=ACCENT,
            padding=(2, 2),
        )
        style.map("Accent.TCheckbutton", foreground=[("active", TEXT_PRIMARY), ("selected", TEXT_PRIMARY)])

        style.configure("Player.TNotebook", background=WINDOW_BG, borderwidth=0, tabmargins=(0, 0, 0, 0))
        style.configure(
            "Player.TNotebook.Tab",
            background=ELEVATED_BG,
            foreground=TEXT_SECONDARY,
            padding=(18, 8),
            borderwidth=0,
            font=("Microsoft YaHei UI", 10, "bold"),
        )
        style.map(
            "Player.TNotebook.Tab",
            background=[("selected", CARD_BG), ("active", SURFACE_BG)],
            foreground=[("selected", TEXT_PRIMARY), ("active", TEXT_PRIMARY)],
        )

        style.configure(
            "Player.Treeview",
            background=TREE_BG,
            fieldbackground=TREE_BG,
            foreground=TEXT_PRIMARY,
            bordercolor=BORDER,
            rowheight=34,
            relief="flat",
            font=("Microsoft YaHei UI", 10),
        )
        style.map("Player.Treeview", background=[("selected", TREE_SELECTED)], foreground=[("selected", TEXT_PRIMARY)])
        style.configure(
            "Player.Treeview.Heading",
            background=CARD_BG,
            foreground=TEXT_SECONDARY,
            borderwidth=0,
            relief="flat",
            font=("Microsoft YaHei UI", 9, "bold"),
            padding=(8, 10),
        )
        style.map("Player.Treeview.Heading", background=[("active", CARD_BG)])

        style.configure(
            "Vertical.TScrollbar",
            background=ELEVATED_BG,
            troughcolor=TREE_BG,
            bordercolor=TREE_BG,
            arrowcolor=TEXT_SECONDARY,
        )

    def _build_ui(self) -> None:
        root = ttk.Frame(self, style="App.TFrame", padding=18)
        root.pack(fill=tk.BOTH, expand=True)
        root.columnconfigure(0, weight=5)
        root.columnconfigure(1, weight=2)
        root.rowconfigure(1, weight=1)

        self._build_hero_panel(root)
        self._build_main_content(root)
        self._build_sidebar(root)

        status = ttk.Label(root, textvariable=self.status_var, anchor="w", style="Status.TLabel")
        status.grid(row=2, column=0, columnspan=2, sticky="ew", pady=(14, 0))

    def _build_hero_panel(self, root: ttk.Frame) -> None:
        hero = tk.Frame(root, bg=SURFACE_BG, highlightthickness=1, highlightbackground=BORDER)
        hero.grid(row=0, column=0, columnspan=2, sticky="ew")
        hero.grid_columnconfigure(0, weight=1)
        hero.grid_columnconfigure(1, weight=1)

        left = tk.Frame(hero, bg=SURFACE_BG, padx=20, pady=18)
        left.grid(row=0, column=0, sticky="nsew")
        ttk.Label(left, text="Bili Music Player", style="HeroTitle.TLabel").pack(anchor="w")
        ttk.Label(
            left,
            text="搜 B 站 MV，抽音频直接听。界面我给你往音乐播放器的方向收了，走在线串流 + 本地缓存双模式。",
            style="HeroSub.TLabel",
            wraplength=560,
            justify="left",
        ).pack(anchor="w", pady=(6, 10))
        ttk.Label(left, textvariable=self.stats_var, style="HeroSub.TLabel").pack(anchor="w")

        search_wrap = tk.Frame(hero, bg=SURFACE_BG, padx=20, pady=18)
        search_wrap.grid(row=0, column=1, sticky="nsew")
        search_wrap.grid_columnconfigure(0, weight=1)

        self.query_entry = ttk.Entry(search_wrap, textvariable=self.query_var, style="Search.TEntry")
        self.query_entry.grid(row=0, column=0, columnspan=3, sticky="ew")
        self.query_entry.bind("<Return>", lambda _event: self._search_or_open_from_query())

        self.search_button = ttk.Button(search_wrap, text="搜索歌曲", style="Primary.TButton", command=self._start_search)
        self.search_button.grid(row=1, column=0, sticky="ew", pady=(12, 0), padx=(0, 8))

        self.play_link_button = ttk.Button(
            search_wrap,
            text="播放链接 / BV",
            style="Secondary.TButton",
            command=self._start_stream_from_query,
        )
        self.play_link_button.grid(row=1, column=1, sticky="ew", pady=(12, 0), padx=(0, 8))

        self.save_link_button = ttk.Button(
            search_wrap,
            text="保存到本地",
            style="Secondary.TButton",
            command=self._start_cache_from_query,
        )
        self.save_link_button.grid(row=1, column=2, sticky="ew", pady=(12, 0))

        self.cache_check = ttk.Checkbutton(
            search_wrap,
            text="在线播放时顺手缓存到本地",
            variable=self.cache_after_stream_var,
            style="Accent.TCheckbutton",
        )
        self.cache_check.grid(row=2, column=0, columnspan=3, sticky="w", pady=(12, 0))

    def _build_main_content(self, root: ttk.Frame) -> None:
        content = ttk.Frame(root, style="App.TFrame")
        content.grid(row=1, column=0, sticky="nsew", pady=(16, 0), padx=(0, 16))
        content.columnconfigure(0, weight=1)
        content.rowconfigure(1, weight=1)
        content.rowconfigure(2, weight=0)

        header = ttk.Frame(content, style="App.TFrame")
        header.grid(row=0, column=0, sticky="ew", pady=(0, 10))
        header.columnconfigure(0, weight=1)

        ttk.Label(header, text="歌单区域", style="SectionTitle.TLabel").grid(row=0, column=0, sticky="w")
        ttk.Label(header, textvariable=self.search_summary_var, style="Summary.TLabel").grid(row=1, column=0, sticky="w", pady=(4, 0))

        self.notebook = ttk.Notebook(content, style="Player.TNotebook")
        self.notebook.grid(row=1, column=0, sticky="nsew")

        search_tab = ttk.Frame(self.notebook, style="Surface.TFrame", padding=14)
        favorites_tab = ttk.Frame(self.notebook, style="Surface.TFrame", padding=14)
        local_tab = ttk.Frame(self.notebook, style="Surface.TFrame", padding=14)
        self.notebook.add(search_tab, text="搜索结果")
        self.notebook.add(favorites_tab, text="我喜欢")
        self.notebook.add(local_tab, text="本地缓存")

        self._build_search_results_tab(search_tab)
        self._build_favorites_tab(favorites_tab)
        self._build_local_library_tab(local_tab)
        self._build_player_panel(content)

    def _build_search_results_tab(self, tab: ttk.Frame) -> None:
        tab.columnconfigure(0, weight=1)
        tab.rowconfigure(0, weight=1)

        columns = ("title", "uploader", "duration", "views")
        self.results_tree = ttk.Treeview(tab, columns=columns, show="headings", selectmode="browse", style="Player.Treeview")
        self.results_tree.heading("title", text="标题")
        self.results_tree.heading("uploader", text="UP 主")
        self.results_tree.heading("duration", text="时长")
        self.results_tree.heading("views", text="播放")
        self.results_tree.column("title", width=520, minwidth=260, stretch=True)
        self.results_tree.column("uploader", width=160, minwidth=100, stretch=False)
        self.results_tree.column("duration", width=80, minwidth=70, stretch=False, anchor="center")
        self.results_tree.column("views", width=90, minwidth=70, stretch=False, anchor="e")
        self.results_tree.grid(row=0, column=0, sticky="nsew")
        self.results_tree.bind("<Double-Button-1>", lambda _event: self._play_selected_result())
        self.results_tree.bind("<<TreeviewSelect>>", lambda _event: self._on_search_result_selected())
        self.results_tree.tag_configure("odd", background=TREE_BG)
        self.results_tree.tag_configure("even", background=TREE_ALT)

        scrollbar = ttk.Scrollbar(tab, orient=tk.VERTICAL, command=self.results_tree.yview)
        scrollbar.grid(row=0, column=1, sticky="ns")
        self.results_tree.configure(yscrollcommand=scrollbar.set)

        actions = ttk.Frame(tab, style="Surface.TFrame")
        actions.grid(row=1, column=0, columnspan=2, sticky="ew", pady=(12, 0))
        actions.columnconfigure(3, weight=1)

        self.play_result_button = ttk.Button(actions, text="立即播放", style="Primary.TButton", command=self._play_selected_result)
        self.play_result_button.grid(row=0, column=0, padx=(0, 8))

        self.save_result_button = ttk.Button(actions, text="缓存选中", style="Secondary.TButton", command=self._cache_selected_result)
        self.save_result_button.grid(row=0, column=1, padx=(0, 8))

        self.favorite_result_button = ttk.Button(actions, text="❤ 加入我喜欢", style="Secondary.TButton", command=self._favorite_selected_result)
        self.favorite_result_button.grid(row=0, column=2, padx=(0, 8))

        ttk.Button(actions, text="清空结果", style="Secondary.TButton", command=self._clear_search_results).grid(row=0, column=3)

    def _build_favorites_tab(self, tab: ttk.Frame) -> None:
        tab.columnconfigure(0, weight=1)
        tab.rowconfigure(0, weight=1)

        columns = ("title", "uploader", "duration", "views")
        self.favorites_tree = ttk.Treeview(tab, columns=columns, show="headings", selectmode="browse", style="Player.Treeview")
        self.favorites_tree.heading("title", text="标题")
        self.favorites_tree.heading("uploader", text="来源")
        self.favorites_tree.heading("duration", text="时长")
        self.favorites_tree.heading("views", text="播放")
        self.favorites_tree.column("title", width=500, minwidth=240, stretch=True)
        self.favorites_tree.column("uploader", width=180, minwidth=100, stretch=False)
        self.favorites_tree.column("duration", width=80, minwidth=70, stretch=False, anchor="center")
        self.favorites_tree.column("views", width=90, minwidth=70, stretch=False, anchor="e")
        self.favorites_tree.grid(row=0, column=0, sticky="nsew")
        self.favorites_tree.bind("<Double-Button-1>", lambda _event: self._play_selected_favorite())
        self.favorites_tree.bind("<<TreeviewSelect>>", lambda _event: self._on_favorite_selected())
        self.favorites_tree.tag_configure("odd", background=TREE_BG)
        self.favorites_tree.tag_configure("even", background=TREE_ALT)

        scrollbar = ttk.Scrollbar(tab, orient=tk.VERTICAL, command=self.favorites_tree.yview)
        scrollbar.grid(row=0, column=1, sticky="ns")
        self.favorites_tree.configure(yscrollcommand=scrollbar.set)

        actions = ttk.Frame(tab, style="Surface.TFrame")
        actions.grid(row=1, column=0, columnspan=2, sticky="ew", pady=(12, 0))
        actions.columnconfigure(5, weight=1)

        self.play_favorite_button = ttk.Button(actions, text="播放喜欢", style="Primary.TButton", command=self._play_selected_favorite)
        self.play_favorite_button.grid(row=0, column=0, padx=(0, 8))
        self.save_favorite_button = ttk.Button(actions, text="缓存到本地", style="Secondary.TButton", command=self._cache_selected_favorite)
        self.save_favorite_button.grid(row=0, column=1, padx=(0, 8))
        self.remove_favorite_button = ttk.Button(actions, text="移出我喜欢", style="Secondary.TButton", command=self._remove_selected_favorite)
        self.remove_favorite_button.grid(row=0, column=2, padx=(0, 8))
        self.favorite_mode_button = ttk.Button(actions, textvariable=self.favorite_mode_var, style="Secondary.TButton", command=self._toggle_favorite_play_mode)
        self.favorite_mode_button.grid(row=0, column=3, padx=(0, 8))
        self.import_favorites_button = ttk.Button(actions, text="导入 CSV 歌单", style="Secondary.TButton", command=self._import_favorites_from_csv)
        self.import_favorites_button.grid(row=0, column=4)

    def _build_local_library_tab(self, tab: ttk.Frame) -> None:
        tab.columnconfigure(0, weight=1)
        tab.rowconfigure(0, weight=1)

        columns = ("title", "duration", "file")
        self.local_tree = ttk.Treeview(tab, columns=columns, show="headings", selectmode="browse", style="Player.Treeview")
        self.local_tree.heading("title", text="标题")
        self.local_tree.heading("duration", text="时长")
        self.local_tree.heading("file", text="文件")
        self.local_tree.column("title", width=480, minwidth=240, stretch=True)
        self.local_tree.column("duration", width=80, minwidth=70, stretch=False, anchor="center")
        self.local_tree.column("file", width=260, minwidth=160, stretch=False)
        self.local_tree.grid(row=0, column=0, sticky="nsew")
        self.local_tree.bind("<Double-Button-1>", lambda _event: self._play_selected_local_track())
        self.local_tree.bind("<<TreeviewSelect>>", lambda _event: self._on_local_track_selected())
        self.local_tree.tag_configure("odd", background=TREE_BG)
        self.local_tree.tag_configure("even", background=TREE_ALT)

        scrollbar = ttk.Scrollbar(tab, orient=tk.VERTICAL, command=self.local_tree.yview)
        scrollbar.grid(row=0, column=1, sticky="ns")
        self.local_tree.configure(yscrollcommand=scrollbar.set)

        actions = ttk.Frame(tab, style="Surface.TFrame")
        actions.grid(row=1, column=0, columnspan=2, sticky="ew", pady=(12, 0))
        actions.columnconfigure(2, weight=1)

        self.play_local_button = ttk.Button(actions, text="播放缓存", style="Primary.TButton", command=self._play_selected_local_track)
        self.play_local_button.grid(row=0, column=0, padx=(0, 8))

        ttk.Button(actions, text="刷新缓存", style="Secondary.TButton", command=self._refresh_library).grid(row=0, column=1)

    def _build_sidebar(self, root: ttk.Frame) -> None:
        sidebar = ttk.Frame(root, style="App.TFrame")
        sidebar.grid(row=1, column=1, sticky="nsew", pady=(16, 0))
        sidebar.columnconfigure(0, weight=1)

        self._build_now_playing_card(sidebar)
        self._build_selection_card(sidebar)
        self._build_favorites_card(sidebar)
        self._build_library_card(sidebar)

    def _build_now_playing_card(self, parent: ttk.Frame) -> None:
        card = tk.Frame(parent, bg=CARD_BG, padx=18, pady=18, highlightthickness=1, highlightbackground=BORDER)
        card.grid(row=0, column=0, sticky="ew")

        ttk.Label(card, text="Now Playing", style="CardTitle.TLabel").pack(anchor="w")
        self.artwork_label = tk.Label(
            card,
            bg=ELEVATED_BG,
            width=280,
            height=280,
            image=self._artwork_placeholder,
            compound="center",
            bd=0,
            highlightthickness=0,
        )
        self.artwork_label.pack(fill="x", pady=(12, 12))
        ttk.Label(card, textvariable=self.now_playing_var, style="NowPlaying.TLabel", wraplength=300, justify="left").pack(anchor="w", pady=(0, 8))
        ttk.Label(card, textvariable=self.play_mode_var, style="Meta.TLabel").pack(anchor="w")
        ttk.Label(card, textvariable=self.transport_hint_var, style="Meta.TLabel", wraplength=300, justify="left").pack(anchor="w", pady=(4, 0))

    def _build_selection_card(self, parent: ttk.Frame) -> None:
        card = tk.Frame(parent, bg=CARD_BG, padx=18, pady=18, highlightthickness=1, highlightbackground=BORDER)
        card.grid(row=1, column=0, sticky="ew", pady=(14, 0))

        ttk.Label(card, text="当前选中", style="CardTitle.TLabel").pack(anchor="w")
        ttk.Label(card, textvariable=self.selection_title_var, style="NowPlaying.TLabel", wraplength=300, justify="left").pack(anchor="w", pady=(10, 8))
        ttk.Label(card, textvariable=self.selection_meta_var, style="Meta.TLabel", wraplength=300, justify="left").pack(anchor="w")
        ttk.Label(card, textvariable=self.selection_extra_var, style="Meta.TLabel", wraplength=300, justify="left").pack(anchor="w", pady=(8, 0))

    def _build_favorites_card(self, parent: ttk.Frame) -> None:
        card = tk.Frame(parent, bg=CARD_BG, padx=18, pady=18, highlightthickness=1, highlightbackground=BORDER)
        card.grid(row=2, column=0, sticky="ew", pady=(14, 0))

        ttk.Label(card, text="我喜欢", style="CardTitle.TLabel").pack(anchor="w")
        ttk.Label(card, textvariable=self.favorite_summary_var, style="Meta.TLabel", wraplength=300, justify="left").pack(anchor="w", pady=(10, 0))
        ttk.Label(card, textvariable=self.favorite_mode_summary_var, style="Meta.TLabel", wraplength=300, justify="left").pack(anchor="w", pady=(6, 0))

    def _build_library_card(self, parent: ttk.Frame) -> None:
        card = tk.Frame(parent, bg=CARD_BG, padx=18, pady=18, highlightthickness=1, highlightbackground=BORDER)
        card.grid(row=3, column=0, sticky="ew", pady=(14, 0))

        ttk.Label(card, text="本地曲库", style="CardTitle.TLabel").pack(anchor="w")
        ttk.Label(card, textvariable=self.library_summary_var, style="Meta.TLabel", wraplength=300, justify="left").pack(anchor="w", pady=(10, 0))

    def _build_player_panel(self, root: ttk.Frame) -> None:
        player_frame = tk.Frame(root, bg=CARD_BG, padx=18, pady=16, highlightthickness=1, highlightbackground=BORDER)
        player_frame.grid(row=2, column=0, sticky="ew", pady=(14, 0))
        player_frame.grid_columnconfigure(0, weight=1)

        top_row = tk.Frame(player_frame, bg=CARD_BG)
        top_row.grid(row=0, column=0, sticky="ew")
        top_row.grid_columnconfigure(0, weight=1)

        ttk.Label(top_row, text="播放控制台", style="CardTitle.TLabel").grid(row=0, column=0, sticky="w")
        ttk.Label(top_row, textvariable=self.time_var, style="Meta.TLabel").grid(row=0, column=1, sticky="e")

        self.progress_scale = ttk.Scale(
            player_frame,
            from_=0,
            to=PROGRESS_SCALE_MAX,
            variable=self.progress_var,
        )
        self.progress_scale.grid(row=1, column=0, sticky="ew", pady=(12, 10))
        self.progress_scale.bind("<ButtonPress-1>", self._begin_progress_drag)
        self.progress_scale.bind("<ButtonRelease-1>", self._finish_progress_drag)

        controls = tk.Frame(player_frame, bg=CARD_BG)
        controls.grid(row=2, column=0, sticky="ew")
        controls.grid_columnconfigure(4, weight=1)

        self.pause_button = ttk.Button(controls, textvariable=self.pause_text, style="Primary.TButton", command=self._pause_or_resume)
        self.pause_button.grid(row=0, column=0, padx=(0, 8))
        self.stop_button = ttk.Button(controls, text="停止", style="Secondary.TButton", command=self._stop)
        self.stop_button.grid(row=0, column=1, padx=(0, 8))
        ttk.Label(controls, text="拖动进度仅对本地缓存稳定生效", style="Meta.TLabel").grid(row=0, column=4, sticky="e")

    def _load_library(self) -> None:
        for child in self.local_tree.get_children():
            self.local_tree.delete(child)
        self.library.prune_missing()
        self.library.save()
        tracks = self.library.tracks
        for index, track in enumerate(tracks):
            self.local_tree.insert(
                "",
                tk.END,
                iid=f"local:{index}",
                tags=("even" if index % 2 else "odd",),
                values=(
                    track.title,
                    self._format_time(track.duration),
                    Path(track.file_path).name,
                ),
            )
        self.library_summary_var.set(
            "还没有本地缓存。先搜一首歌试试。" if not tracks else f"已经缓存 {len(tracks)} 首，可以双击离线播放。"
        )

    def _load_favorites(self) -> None:
        for child in self.favorites_tree.get_children():
            self.favorites_tree.delete(child)
        favorites = self.library.favorites
        for index, favorite in enumerate(favorites):
            self.favorites_tree.insert(
                "",
                tk.END,
                iid=f"favorite:{index}",
                tags=("even" if index % 2 else "odd",),
                values=(
                    favorite.title,
                    favorite.uploader or "",
                    self._format_time(favorite.duration),
                    self._format_count(favorite.view_count),
                ),
            )
        self.favorite_summary_var.set(
            "把喜欢的歌收进列表，下次一点就播。" if not favorites else f"已收藏 {len(favorites)} 首喜欢歌曲，双击就能播。"
        )

    def _refresh_library(self) -> None:
        self._load_library()
        self._load_favorites()
        self._refresh_dashboard()
        self._set_status("本地缓存已刷新。")

    def _start_search(self) -> None:
        if self.is_search_running:
            return

        keyword = self.query_var.get().strip()
        if not keyword:
            self._show_error("请输入搜索关键词。")
            return

        self.is_search_running = True
        self._set_search_controls_state(tk.DISABLED)
        self.search_summary_var.set(f"正在搜索 “{keyword}” ...")
        self._set_status(f"正在搜索：{keyword}")
        thread = threading.Thread(target=self._search_worker, args=(keyword,), daemon=True)
        thread.start()

    def _search_worker(self, keyword: str) -> None:
        try:
            results = self.downloader.search(
                keyword,
                SEARCH_LIMIT,
                lambda message: self.events.put(("status", message)),
            )
            self.events.put(("search_done", results))
        except AudioDownloadError as exc:
            self.events.put(("search_error", str(exc)))
        except Exception as exc:
            self.events.put(("search_error", f"发生未知错误：{exc}"))

    def _search_or_open_from_query(self) -> None:
        query = self.query_var.get().strip()
        if self.downloader.is_video_reference(query):
            self._start_stream_from_query()
        else:
            self._start_search()

    def _start_stream_from_query(self) -> None:
        url = self._query_video_url_or_error()
        if url is not None:
            self.current_favorite_url = None
            self._start_stream_url(url)

    def _start_cache_from_query(self) -> None:
        url = self._query_video_url_or_error()
        if url is not None:
            self._start_cache_url(url, play_after=False)

    def _play_selected_result(self) -> None:
        result = self._selected_search_result()
        if result is not None:
            self.current_favorite_url = None
            self._start_stream_url(result.url)

    def _cache_selected_result(self) -> None:
        result = self._selected_search_result()
        if result is not None:
            self._start_cache_url(result.url, play_after=False)

    def _favorite_selected_result(self) -> None:
        result = self._selected_search_result()
        if result is None:
            return
        self.library.add_favorite(
            FavoriteTrack(
                title=result.title,
                url=result.url,
                duration=result.duration,
                uploader=result.uploader,
                view_count=result.view_count,
                thumbnail_url=result.thumbnail_url,
            )
        )
        self._load_favorites()
        self._refresh_dashboard()
        self._set_status(f"已加入我喜欢：{result.title}")

    def _toggle_favorite_play_mode(self) -> None:
        self.favorite_play_mode = "random" if self.favorite_play_mode == "sequential" else "sequential"
        mode_text = "随机" if self.favorite_play_mode == "random" else "顺序"
        self.favorite_mode_var.set(f"播放模式：{mode_text}")
        self.favorite_mode_summary_var.set(f"当前模式：{mode_text}播放")
        self._set_status(f"我喜欢列表已切换为{mode_text}播放。")

    def _import_favorites_from_csv(self) -> None:
        csv_path = filedialog.askopenfilename(
            title="选择歌单文件",
            filetypes=[("歌单文件", "*.csv;*.xlsx;*.xlsm"), ("CSV 文件", "*.csv"), ("Excel 文件", "*.xlsx;*.xlsm"), ("所有文件", "*.*")],
        )
        if not csv_path:
            return
        self._begin_media_task("正在导入 CSV 歌单...")
        thread = threading.Thread(target=self._import_favorites_worker, args=(csv_path,), daemon=True)
        thread.start()

    def _import_favorites_worker(self, csv_path: str) -> None:
        try:
            imported, skipped = self._read_and_import_csv(
                Path(csv_path),
                progress_callback=lambda message: self.events.put(("status", message)),
            )
            self.events.put(("favorites_imported", (imported, skipped)))
        except Exception as exc:
            self.events.put(("favorites_import_error", str(exc)))

    def _start_stream_url(self, url: str) -> None:
        if self.is_media_task_running:
            return

        self._begin_media_task("正在解析音频流...")
        cache_after_stream = self.cache_after_stream_var.get()
        thread = threading.Thread(target=self._stream_worker, args=(url, cache_after_stream), daemon=True)
        thread.start()

    def _start_cache_url(self, url: str, play_after: bool) -> None:
        if self.is_media_task_running:
            return

        self._begin_media_task("正在保存音频到本地...")
        self._start_cache_worker(url, play_after=play_after)

    def _stream_worker(self, url: str, cache_after_stream: bool) -> None:
        try:
            stream = self.downloader.resolve_stream(url, lambda message: self.events.put(("status", message)))
            self.events.put(("stream_resolved", (stream, cache_after_stream)))
        except AudioDownloadError as exc:
            self.events.put(("media_error", str(exc)))
        except Exception as exc:
            self.events.put(("media_error", f"发生未知错误：{exc}"))

    def _start_cache_worker(self, url: str, play_after: bool) -> None:
        thread = threading.Thread(target=self._cache_worker, args=(url, play_after), daemon=True)
        thread.start()

    def _cache_worker(self, url: str, play_after: bool) -> None:
        try:
            track = self.downloader.download(url, lambda message: self.events.put(("status", message)))
            self.events.put(("cached", (track, play_after)))
        except AudioDownloadError as exc:
            self.events.put(("media_error", str(exc)))
        except Exception as exc:
            self.events.put(("media_error", f"发生未知错误：{exc}"))

    def _start_temp_play_worker(self, url: str) -> None:
        thread = threading.Thread(target=self._temp_play_worker, args=(url,), daemon=True)
        thread.start()

    def _temp_play_worker(self, url: str) -> None:
        try:
            self.events.put(("status", "正在建立临时音频缓存..."))
            track = self.downloader.download_temporary(url, lambda message: self.events.put(("status", message)))
            self.events.put(("temp_ready", track))
        except AudioDownloadError as exc:
            self.events.put(("media_error", str(exc)))
        except Exception as exc:
            self.events.put(("media_error", f"发生未知错误：{exc}"))

    def _process_events(self) -> None:
        while True:
            try:
                event, payload = self.events.get_nowait()
            except queue.Empty:
                break

            if event == "status":
                self._set_status(str(payload))
            elif event == "search_done":
                self._finish_search()
                if isinstance(payload, list):
                    self._show_search_results(payload)
            elif event == "search_error":
                self._finish_search()
                self.search_summary_var.set("搜索失败了，换个关键词或直接贴 BV / 链接试试。")
                self._show_error(str(payload))
            elif event == "stream_resolved":
                stream, cache_after_stream = payload
                if isinstance(stream, AudioStream):
                    self._handle_stream_resolved(stream, bool(cache_after_stream))
            elif event == "temp_ready" and isinstance(payload, Track):
                if self._play_track(payload):
                    self.play_mode_var.set("临时缓存播放中")
                    self.transport_hint_var.set("当前是临时缓存文件，避免了整段视频下载。")
                    self._set_status(f"正在播放临时缓存：{payload.title}")
                self._finish_media_task()
            elif event == "cached":
                track, play_after = payload
                if isinstance(track, Track):
                    self.library.add_or_update(track)
                    self._load_library()
                    self._refresh_dashboard()
                    if play_after:
                        if self._play_track(track):
                            self.play_mode_var.set("已保存后播放")
                            self.transport_hint_var.set("当前播放的是本地缓存，支持拖动进度。")
                            self._set_status(f"已保存并播放：{track.title}")
                    else:
                        self._set_status(f"已保存到本地：{track.title}")
                self._finish_media_task()
            elif event == "favorites_imported":
                imported, skipped = payload
                self._finish_media_task()
                self._load_favorites()
                self._refresh_dashboard()
                self._set_status(f"歌单导入完成：新增 {imported} 首，跳过 {skipped} 首。")
            elif event == "favorites_import_error":
                self._finish_media_task()
                self._show_error(f"导入 CSV 失败：{payload}")
            elif event == "media_error":
                self._finish_media_task()
                self.play_mode_var.set("任务失败")
                self.transport_hint_var.set("这次没拉起来，通常换个结果或稍后再试就行。")
                self._show_error(str(payload))

        self.after(150, self._process_events)

    def _finish_search(self) -> None:
        self.is_search_running = False
        self._set_search_controls_state(tk.NORMAL)

    def _show_search_results(self, results: list[BiliSearchResult]) -> None:
        self.search_results = results
        for child in self.results_tree.get_children():
            self.results_tree.delete(child)
        for index, result in enumerate(results):
            self.results_tree.insert(
                "",
                tk.END,
                iid=f"result:{index}",
                tags=("even" if index % 2 else "odd",),
                values=(
                    result.title,
                    result.uploader or "",
                    self._format_time(result.duration),
                    self._format_count(result.view_count),
                ),
            )
        self.notebook.select(0)
        if results:
            self.results_tree.selection_set("result:0")
            self.results_tree.focus("result:0")
            self._update_search_selection(results[0])
        else:
            self.selection_title_var.set("没有找到结果")
            self.selection_meta_var.set("换个关键词试试，或者直接输入 BV / 链接。")
            self.selection_extra_var.set("也可能是接口暂时没返回可用数据。")
            self._set_artwork_placeholder("暂无封面")
        self.search_summary_var.set(f"找到 {len(results)} 条结果，双击就能播。")
        self._refresh_dashboard()
        self._set_status(f"找到 {len(results)} 条结果，双击或选择后播放。")

    def _clear_search_results(self) -> None:
        self.search_results = []
        for child in self.results_tree.get_children():
            self.results_tree.delete(child)
        self.search_summary_var.set("搜索结果已清空。")
        self.selection_title_var.set("等待选择内容")
        self.selection_meta_var.set("先搜一首歌，我会把结果摆这里。")
        self.selection_extra_var.set("支持在线播放和保存到本地。")
        self._set_artwork_placeholder("等待选歌")
        self._refresh_dashboard()
        self._set_status("搜索结果已清空。")

    def _handle_stream_resolved(self, stream: AudioStream, cache_after_stream: bool) -> None:
        try:
            self.player.play(stream.stream_url)
        except AudioPlayerError:
            fallback = "本地缓存" if cache_after_stream else "临时音频缓存"
            self.play_mode_var.set("远程串流失败，切换兜底方案")
            self.transport_hint_var.set(f"播放器无法直开远程流，正在改用{fallback}。")
            self._set_status(f"播放器无法直接打开远程音频流，改用{fallback}...")
            if cache_after_stream:
                self._start_cache_worker(stream.url, play_after=True)
            else:
                self._start_temp_play_worker(stream.url)
            return

        self.pause_text.set("暂停")
        self._set_now_playing(stream.title, stream.duration)
        self.play_mode_var.set("在线播放中")
        self.transport_hint_var.set("当前优先使用独立音频流，没有先下载完整视频。")
        if cache_after_stream:
            self._set_status(f"正在串流播放：{stream.title}；后台保存中...")
            self._start_cache_worker(stream.url, play_after=False)
        else:
            self._set_status(f"正在串流播放：{stream.title}")
            self._finish_media_task()

    def _selected_search_result(self) -> BiliSearchResult | None:
        selection = self.results_tree.selection()
        if not selection:
            self._show_error("请先在搜索结果中选择一个视频。")
            return None

        index = self._index_from_iid(selection[0], "result:")
        if index is None or index >= len(self.search_results):
            self._show_error("选中的搜索结果无效，请重新搜索。")
            return None
        return self.search_results[index]

    def _selected_local_track(self) -> Track | None:
        selection = self.local_tree.selection()
        if not selection:
            self._show_error("请先在本地缓存中选择一首音频。")
            return None

        tracks = self.library.tracks
        index = self._index_from_iid(selection[0], "local:")
        if index is None or index >= len(tracks):
            self._show_error("选中的缓存条目无效，请刷新缓存。")
            return None
        return tracks[index]

    def _selected_favorite(self) -> FavoriteTrack | None:
        selection = self.favorites_tree.selection()
        if not selection:
            self._show_error("请先在“我喜欢”中选择一首歌。")
            return None

        favorites = self.library.favorites
        index = self._index_from_iid(selection[0], "favorite:")
        if index is None or index >= len(favorites):
            self._show_error("选中的喜欢条目无效，请刷新后重试。")
            return None
        return favorites[index]

    @staticmethod
    def _index_from_iid(iid: str, prefix: str) -> int | None:
        if not iid.startswith(prefix):
            return None
        try:
            return int(iid[len(prefix) :])
        except ValueError:
            return None

    def _on_search_result_selected(self) -> None:
        result = self._selected_search_result_silent()
        if result is not None:
            self._update_search_selection(result)

    def _on_favorite_selected(self) -> None:
        favorite = self._selected_favorite_silent()
        if favorite is not None:
            self.selection_title_var.set(favorite.title)
            self.selection_meta_var.set(
                f"我喜欢 · {favorite.uploader or '未知来源'} · {self._format_time(favorite.duration)}"
            )
            self.selection_extra_var.set("双击直接播放；也可以随时缓存到本地。")
            self._load_artwork_async(favorite.thumbnail_url, favorite.title)

    def _on_local_track_selected(self) -> None:
        track = self._selected_local_track_silent()
        if track is not None:
            self.selection_title_var.set(track.title)
            self.selection_meta_var.set(f"本地缓存 · {self._format_time(track.duration)} · {Path(track.file_path).name}")
            self.selection_extra_var.set("双击即可离线播放；本地文件支持拖动进度。")
            self._set_artwork_placeholder("本地缓存")

    def _selected_search_result_silent(self) -> BiliSearchResult | None:
        selection = self.results_tree.selection()
        if not selection:
            return None
        index = self._index_from_iid(selection[0], "result:")
        if index is None or index >= len(self.search_results):
            return None
        return self.search_results[index]

    def _selected_local_track_silent(self) -> Track | None:
        selection = self.local_tree.selection()
        if not selection:
            return None
        tracks = self.library.tracks
        index = self._index_from_iid(selection[0], "local:")
        if index is None or index >= len(tracks):
            return None
        return tracks[index]

    def _selected_favorite_silent(self) -> FavoriteTrack | None:
        selection = self.favorites_tree.selection()
        if not selection:
            return None
        favorites = self.library.favorites
        index = self._index_from_iid(selection[0], "favorite:")
        if index is None or index >= len(favorites):
            return None
        return favorites[index]

    def _update_search_selection(self, result: BiliSearchResult) -> None:
        uploader = result.uploader or "未知 UP 主"
        views = self._format_count(result.view_count) or "播放量未知"
        duration = self._format_time(result.duration)
        quality_hint = self._infer_quality_hint(result)
        self.selection_title_var.set(result.title)
        self.selection_meta_var.set(f"{uploader} · {duration} · {views} 次播放")
        self.selection_extra_var.set(f"{quality_hint}；双击直接播放，想收藏的话点“缓存选中”。")
        self._load_artwork_async(result.thumbnail_url, result.title)

    def _play_selected_local_track(self) -> None:
        track = self._selected_local_track()
        if track is not None:
            self.current_favorite_url = None
            self._play_track(track)

    def _play_selected_favorite(self) -> None:
        favorite = self._selected_favorite()
        if favorite is not None:
            self.current_favorite_url = favorite.url
            self._start_stream_url(favorite.url)

    def _cache_selected_favorite(self) -> None:
        favorite = self._selected_favorite()
        if favorite is not None:
            self._start_cache_url(favorite.url, play_after=False)

    def _remove_selected_favorite(self) -> None:
        favorite = self._selected_favorite()
        if favorite is None:
            return
        self.library.remove_favorite(favorite.url)
        if self.current_favorite_url == favorite.url:
            self.current_favorite_url = None
        self._load_favorites()
        self._refresh_dashboard()
        self._set_status(f"已移出我喜欢：{favorite.title}")

    def _play_track(self, track: Track) -> bool:
        try:
            self.player.play(track.file_path)
        except AudioPlayerError as exc:
            self._show_error(str(exc))
            return False
        self.pause_text.set("暂停")
        self._set_now_playing(track.title, track.duration)
        self.play_mode_var.set("本地缓存播放中")
        self.transport_hint_var.set("当前播放的是本地文件，支持拖动进度。")
        self._set_status(f"正在播放：{track.title}")
        return True

    def _pause_or_resume(self) -> None:
        try:
            paused = self.player.pause_or_resume()
        except AudioPlayerError as exc:
            self._show_error(str(exc))
            return
        self.pause_text.set("继续" if paused else "暂停")
        self.play_mode_var.set("已暂停" if paused else "继续播放")
        self._set_status("已暂停。" if paused else "继续播放。")

    def _stop(self) -> None:
        self.player.stop()
        self.current_favorite_url = None
        self.pause_text.set("暂停")
        self.play_mode_var.set("已停止")
        self.transport_hint_var.set("在线播放优先，失败时自动转临时缓存。")
        self._clear_now_playing()
        self._set_status("已停止。")

    def _begin_progress_drag(self, _event: tk.Event) -> None:
        self.is_dragging_progress = True

    def _finish_progress_drag(self, _event: tk.Event) -> None:
        if not self.is_dragging_progress:
            return
        self.is_dragging_progress = False

        if self.current_duration is None:
            self._set_status("当前音频没有可用总时长，无法拖动进度。")
            self._sync_progress_to_player()
            return

        target = (self.progress_var.get() / PROGRESS_SCALE_MAX) * self.current_duration
        try:
            self.player.seek(target)
        except AudioPlayerError as exc:
            self._set_status(str(exc))
            self._sync_progress_to_player()
            return

        self.playback_started_at = time.monotonic()
        self._sync_progress_to_player()
        self._set_status(f"已跳转到 {self._format_time(target)}。")

    def _update_progress_ui(self) -> None:
        if self.player.has_finished() and time.monotonic() - self.playback_started_at > 1.0:
            if self.current_duration is not None:
                self.progress_var.set(PROGRESS_SCALE_MAX)
                total_time = self._format_time(self.current_duration)
                self.time_var.set(f"{total_time} / {total_time}")
            self.player.clear_finished()
            self.pause_text.set("暂停")
            if self._try_play_next_favorite():
                return
            self.play_mode_var.set("播放结束")
            self.transport_hint_var.set("可以继续搜下一首，或者从本地缓存里再听一遍。")
            self._set_status("播放结束。")
        elif self.player.current_source is not None:
            self._sync_progress_to_player()

        self.after(500, self._update_progress_ui)

    def _sync_progress_to_player(self) -> None:
        position = self.player.position_seconds()
        if self.current_duration is not None:
            position = min(position, self.current_duration)
            if not self.is_dragging_progress:
                progress = (position / self.current_duration) * PROGRESS_SCALE_MAX if self.current_duration else 0
                self.progress_var.set(progress)
            self.time_var.set(f"{self._format_time(position)} / {self._format_time(self.current_duration)}")
        else:
            if not self.is_dragging_progress:
                self.progress_var.set(0.0)
            self.time_var.set(f"{self._format_time(position)} / --:--")

    def _set_now_playing(self, title: str, duration: float | None) -> None:
                self.current_duration = duration
        self.playback_started_at = time.monotonic()
        self.now_playing_var.set(title)
        self.progress_var.set(0.0)
        self.time_var.set(f"00:00 / {self._format_time(duration)}")

    def _clear_now_playing(self) -> None:
                self.current_duration = None
        self.is_dragging_progress = False
        self.now_playing_var.set("未播放")
        self.progress_var.set(0.0)
        self.time_var.set("00:00 / --:--")
        self._set_artwork_placeholder("未在播放")

    def _query_video_url_or_error(self) -> str | None:
        value = self.query_var.get().strip()
        if not value:
            self._show_error("请输入 Bilibili 视频链接或 BV 号。")
            return None
        if not self.downloader.is_video_reference(value):
            self._show_error("当前输入不是视频链接或 BV 号。关键词请使用“搜索”。")
            return None
        return self.downloader.normalized_video_url(value)

    def _try_play_next_favorite(self) -> bool:
        next_favorite = self._next_favorite_to_play()
        if next_favorite is None:
            self.current_favorite_url = None
            return False

        self.current_favorite_url = next_favorite.url
        self.play_mode_var.set("我喜欢列表继续播放")
        self.transport_hint_var.set(
            f"当前按{'随机' if self.favorite_play_mode == 'random' else '顺序'}模式继续播放我喜欢列表。"
        )
        self._set_status(f"继续播放我喜欢：{next_favorite.title}")
        self._start_stream_url(next_favorite.url)
        return True

    def _next_favorite_to_play(self) -> FavoriteTrack | None:
        favorites = self.library.favorites
        if not favorites or self.current_favorite_url is None:
            return None
        if len(favorites) == 1:
            return favorites[0]

        if self.favorite_play_mode == "random":
            candidates = [favorite for favorite in favorites if favorite.url != self.current_favorite_url]
            return random.choice(candidates or favorites)

        current_index = next(
            (index for index, favorite in enumerate(favorites) if favorite.url == self.current_favorite_url),
            None,
        )
        if current_index is None:
            return favorites[0]
        return favorites[(current_index + 1) % len(favorites)]

    def _read_and_import_csv(self, csv_path: Path, progress_callback=None) -> tuple[int, int]:
        imported = 0
        skipped = 0

        rows = self._load_playlist_rows(csv_path)
        if not rows:
            raise ValueError("歌单里没有可读取的内容。")

        total = len(rows)
        for index, row in enumerate(rows, start=1):
            if progress_callback is not None:
                progress_callback(f"正在导入歌单：第 {index}/{total} 首...")

            favorite = self._favorite_from_csv_row(row, progress_callback=progress_callback)
            if favorite is None:
                skipped += 1
                continue
            self.library.add_favorite(favorite)
            imported += 1

            if index < total:
                delay_seconds = random.uniform(*IMPORT_DELAY_RANGE)
                if progress_callback is not None:
                    progress_callback(f"已导入 {imported} 首，暂停 {delay_seconds:.1f} 秒后继续...")
                time.sleep(delay_seconds)

        self._save_search_cache()
        return imported, skipped

    def _load_playlist_rows(self, csv_path: Path) -> list[dict[str, str]]:
        suffix = csv_path.suffix.casefold()
        header = csv_path.read_bytes()[:4]
        if suffix in {".xlsx", ".xlsm"} or header.startswith(b"PK"):
            return self._load_excel_rows(csv_path)
        return self._load_text_csv_rows(csv_path)

    def _load_text_csv_rows(self, csv_path: Path) -> list[dict[str, str]]:
        last_error: Exception | None = None
        for encoding in ("utf-8-sig", "utf-8", "gb18030", "gbk", "utf-16"):
            try:
                with csv_path.open("r", encoding=encoding, newline="") as handle:
                    reader = csv.DictReader(handle)
                    rows = [dict(row) for row in reader if row is not None]
                if rows:
                    return rows
            except Exception as exc:
                last_error = exc
                continue
        raise ValueError(f"无法识别 CSV 编码，请尝试另存为 UTF-8 / CSV。原始错误：{last_error}")

    def _load_excel_rows(self, csv_path: Path) -> list[dict[str, str]]:
        if load_workbook is None:
            raise ValueError("当前环境缺少 openpyxl，无法读取 Excel 格式歌单。")

        workbook = load_workbook(BytesIO(csv_path.read_bytes()), read_only=True, data_only=True)
        sheet = workbook.active
        rows_iter = sheet.iter_rows(values_only=True)
        try:
            header_row = next(rows_iter)
        except StopIteration:
            return []

        headers = [str(cell).strip() if cell is not None else "" for cell in header_row]
        normalized_headers = [header if header else f"column_{index}" for index, header in enumerate(headers)]
        rows: list[dict[str, str]] = []
        for values in rows_iter:
            row = {
                normalized_headers[index]: ("" if value is None else str(value).strip())
                for index, value in enumerate(values)
                if index < len(normalized_headers)
            }
            if any(value for value in row.values()):
                rows.append(row)
        return rows

    def _favorite_from_csv_row(self, row: dict[str, str], progress_callback=None) -> FavoriteTrack | None:
        url_value = self._csv_value(row, "url", "链接", "link", "网址", "video_url")
        bv_value = self._csv_value(row, "bv", "bvid", "BV号", "bv号")
        title_value = self._csv_value(row, "title", "歌名", "歌曲", "name", "keyword", "关键词")
        uploader_value = self._csv_value(row, "uploader", "up主", "歌手", "artist", "作者")

        if bv_value:
            url_value = bv_value

        if url_value:
            normalized_url = self.downloader.normalized_video_url(url_value)
            return FavoriteTrack(
                title=title_value or normalized_url,
                url=normalized_url,
                uploader=uploader_value,
            )

        if not title_value:
            return None

        search_query = " ".join(part for part in (title_value, uploader_value) if part)
        cache_key = self._normalize_cache_key(search_query)
        cached = self.search_cache.get(cache_key)
        if isinstance(cached, dict) and cached.get("url"):
            return FavoriteTrack(
                title=str(cached.get("title") or title_value),
                url=str(cached["url"]),
                duration=cached.get("duration"),
                uploader=cached.get("uploader"),
                view_count=cached.get("view_count"),
                thumbnail_url=cached.get("thumbnail_url"),
            )

        for retry_index in range(len(IMPORT_RETRY_DELAYS) + 1):
            try:
                results = self.downloader.search(search_query, 1)
                if not results:
                    return None
                result = results[0]
                self.search_cache[cache_key] = {
                    "title": result.title,
                    "url": result.url,
                    "duration": result.duration,
                    "uploader": result.uploader,
                    "view_count": result.view_count,
                    "thumbnail_url": result.thumbnail_url,
                }
                return FavoriteTrack(
                    title=result.title,
                    url=result.url,
                    duration=result.duration,
                    uploader=result.uploader,
                    view_count=result.view_count,
                    thumbnail_url=result.thumbnail_url,
                )
            except AudioDownloadError as exc:
                message = str(exc)
                if "412" not in message or retry_index >= len(IMPORT_RETRY_DELAYS):
                    return None
                delay_seconds = IMPORT_RETRY_DELAYS[retry_index]
                if progress_callback is not None:
                    progress_callback(
                        f"B站返回 412，正在等待 {delay_seconds:.0f} 秒后重试：{search_query}"
                    )
                time.sleep(delay_seconds)

        return None

    @staticmethod
    def _csv_value(row: dict[str, str], *keys: str) -> str | None:
        normalized = {str(key).strip().casefold(): value for key, value in row.items() if key is not None}
        for key in keys:
            value = normalized.get(key.casefold())
            if isinstance(value, str) and value.strip():
                return value.strip()
        return None

    def _load_search_cache(self) -> dict[str, dict]:
        try:
            if SEARCH_CACHE_FILE.exists():
                return json.loads(SEARCH_CACHE_FILE.read_text(encoding="utf-8"))
        except Exception:
            pass
        return {}

    def _save_search_cache(self) -> None:
        try:
            SEARCH_CACHE_FILE.write_text(json.dumps(self.search_cache, ensure_ascii=False, indent=2), encoding="utf-8")
        except Exception:
            pass

    @staticmethod
    def _normalize_cache_key(value: str) -> str:
        return " ".join(value.casefold().split())

    def _begin_media_task(self, message: str) -> None:
        self.is_media_task_running = True
        self._set_media_controls_state(tk.DISABLED)
        self.play_mode_var.set("处理中")
        self.transport_hint_var.set("正在向 B 站拉取可播放音频流，请等一下。")
        self._set_status(message)

    def _finish_media_task(self) -> None:
        self.is_media_task_running = False
        self._set_media_controls_state(tk.NORMAL)

    def _set_search_controls_state(self, state: str) -> None:
        self.search_button.configure(state=state)

    def _set_media_controls_state(self, state: str) -> None:
        for widget in (
            self.play_link_button,
            self.save_link_button,
            self.play_result_button,
            self.save_result_button,
            self.favorite_result_button,
            self.play_favorite_button,
            self.save_favorite_button,
            self.remove_favorite_button,
            self.favorite_mode_button,
            self.import_favorites_button,
            self.play_local_button,
            self.cache_check,
        ):
            widget.configure(state=state)

    def _show_error(self, message: str) -> None:
        self._set_status(message)
        messagebox.showerror("Bili Music Player", message)

    def _set_status(self, message: str) -> None:
        self.status_var.set(message)

    def _refresh_dashboard(self) -> None:
        self.stats_var.set(
            f"{len(self.library.favorites)} 首喜欢 · {len(self.library.tracks)} 首缓存 · {len(self.search_results)} 条搜索结果"
        )
        self.favorite_mode_summary_var.set(
            f"当前模式：{'随机' if self.favorite_play_mode == 'random' else '顺序'}播放"
        )

    def _infer_quality_hint(self, result: BiliSearchResult) -> str:
        title = result.title.casefold()
        if "官方" in result.title or "mv" in title or "music video" in title:
            return "这条更像官方 MV / 正片"
        if "live" in title or "现场" in result.title:
            return "这条像现场版，音频可能不是录音室版本"
        if "cover" in title or "翻唱" in result.title:
            return "这条像翻唱版本"
        return "这条看起来适合直接拿来听"

    def _create_placeholder_artwork(self, text: str = "Bili\nMusic"):
        if Image is None or ImageTk is None or ImageDraw is None:
            image = tk.PhotoImage(width=280, height=280)
            image.put(ELEVATED_BG, to=(0, 0, 280, 280))
            return image

        image = Image.new("RGB", (560, 560), "#1b2130")
        draw = ImageDraw.Draw(image)
        for index, color in enumerate(("#ff5c7c", "#ff7a95", "#6e8bff")):
            offset = 42 + index * 26
            draw.rounded_rectangle((offset, offset, 560 - offset, 560 - offset), radius=48, outline=color, width=8)
        draw.text((280, 280), text, anchor="mm", fill="#f7f8fb")
        return ImageTk.PhotoImage(image.resize((280, 280)))

    def _set_artwork_placeholder(self, text: str) -> None:
        placeholder = self._create_placeholder_artwork(text)
                self.artwork_label.configure(image=placeholder)
        self.artwork_label.image = placeholder

    def _load_artwork_async(self, thumbnail_url: str | None, fallback_text: str) -> None:
        self._thumbnail_request_token += 1
        token = self._thumbnail_request_token
        self._set_artwork_placeholder("封面加载中")
        if not thumbnail_url or Image is None or ImageTk is None:
            self._set_artwork_placeholder(fallback_text[:10] or "暂无封面")
            return

        thread = threading.Thread(
            target=self._thumbnail_worker,
            args=(thumbnail_url, fallback_text, token),
            daemon=True,
        )
        thread.start()

    def _thumbnail_worker(self, thumbnail_url: str, fallback_text: str, token: int) -> None:
        image = self._fetch_thumbnail_image(thumbnail_url)
        self.after(0, lambda: self._apply_thumbnail_image(image, fallback_text, token))

    def _fetch_thumbnail_image(self, thumbnail_url: str):
        if Image is None or ImageTk is None or ImageOps is None:
            return None
        try:
            request = urllib.request.Request(thumbnail_url, headers={"User-Agent": "Mozilla/5.0"})
            with urllib.request.urlopen(request, timeout=15) as response:
                raw = response.read()
            image = Image.open(BytesIO(raw)).convert("RGB")
            return ImageOps.fit(image, (280, 280), method=Image.Resampling.LANCZOS)
        except Exception:
            return None

    def _apply_thumbnail_image(self, image, fallback_text: str, token: int) -> None:
        if token != self._thumbnail_request_token:
            return
        if image is None or ImageTk is None:
            self._set_artwork_placeholder(fallback_text[:10] or "暂无封面")
            return
        artwork = ImageTk.PhotoImage(image)
                self.artwork_label.configure(image=artwork)
        self.artwork_label.image = artwork

    @staticmethod
    def _format_time(seconds: float | None) -> str:
        if seconds is None:
            return "--:--"

        total_seconds = max(0, int(seconds))
        minutes, second = divmod(total_seconds, 60)
        hour, minute = divmod(minutes, 60)
        if hour:
            return f"{hour}:{minute:02d}:{second:02d}"
        return f"{minute:02d}:{second:02d}"

    @staticmethod
    def _format_count(value: int | None) -> str:
        return f"{value:,}" if value is not None else ""

    def _on_close(self) -> None:
        self.player.close()
        self.downloader.cleanup_temporary_files()
        self.destroy()


def main() -> None:
    app = BiliMusicPlayerApp()
    app.mainloop()
